#!/usr/bin/env python3
"""Build the final Week 1 deck from the editable UPDATED PowerPoint.

Requires python-pptx==1.0.2, PyMuPDF==1.26.5, and LibreOffice for PDF export.
Run from any directory; assets and source paths are relative to this script.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unicodedata
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt


SLIDES = Path(__file__).resolve().parent
ROOT = SLIDES.parent
SOURCE = ROOT / "ROS_BASICS_UPDATED.pptx"
OUTPUT = ROOT / "ROS_BASICS_FINAL.pptx"
ASSETS = SLIDES / "assets"
QOS_REFERENCE = (
    "https://github.com/ros2/ros2_documentation/blob/jazzy/source/Concepts/"
    "Intermediate/About-Quality-of-Service-Settings.rst"
)
RMW_REFERENCE = (
    "https://github.com/ros2/ros2_documentation/blob/jazzy/source/How-To-Guides/"
    "Working-with-multiple-RMW-implementations.rst"
)
COLORS = {
    "bg": "F7F9FC", "ink": "15283B", "muted": "526477",
    "blue": "2459D3", "teal": "137F79", "paleBlue": "EAF0FC",
    "paleTeal": "E5F3F0", "white": "FFFFFF", "line": "D7E0E9",
}
TITLES = [
    "Robotics &\nAutonomous Systems", "About me", "ROS basics",
    "What is ROS, and why do we need it?", "What is ROS used for?",
    "ROS 1 → ROS 2", "Nodes and the ROS graph", "Topics carry typed messages",
    "Choose the right interface", "ROS domain",
    "ROS 2 communication behind the scenes", "QoS shapes delivery",
    "Lifecycle of a node", "Inspect before debugging", "Demo time",
    "DDS discovery and RTPS", "QoS compatibility", "Executors and concurrency",
    "Managed lifecycle transitions", "Migrating from ROS 1",
]
MINUTES = [0, 1, .5, 2, 1.5, 2, 2, 2, 2, 2, 3, 2, 3, 2, 20, 0, 0, 0, 0, 0]


def rgb(name):
    return RGBColor.from_string(COLORS.get(name, name))


def text(slide, x, y, w, h, value, size=26, color="ink", bold=False,
         align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.0
        for run in p.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
    return shape


def box(slide, x, y, w, h, fill="white"):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)
    )
    shape.adjustments[0] = .09
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    return shape


def find(slide, value):
    return next(s for s in slide.shapes if s.has_text_frame and s.text == value)


def replace(shape, value, size=None):
    """Replace text while retaining the source run and paragraph formatting."""
    frame = shape.text_frame
    p = frame.paragraphs[0]
    rpr = deepcopy(p.runs[0]._r.rPr) if p.runs else None
    ppr = deepcopy(p._p.pPr)
    frame.clear()
    for index, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line
        if ppr is not None:
            old = p._p.pPr
            if old is not None:
                p._p.remove(old)
            p._p.insert(0, deepcopy(ppr))
        if p.runs and rpr is not None:
            p.runs[0]._r.insert(0, deepcopy(rpr))
        if size is not None:
            for run in p.runs:
                run.font.size = Pt(size)


def remove(shape):
    element = shape._element
    element.getparent().remove(element)


def clear_body(slide):
    for shape in list(slide.shapes):
        if Pt(190) <= shape.top < Pt(490):
            remove(shape)


def new_slide(prs, template, title, subtitle):
    slide = prs.slides.add_slide(template.slide_layout)
    for shape in list(slide.shapes):
        remove(shape)
    slide._element.cSld.insert(0, deepcopy(template._element.cSld.bg))
    for shape in template.shapes:
        if shape.top < Pt(190) or shape.top >= Pt(490):
            slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
    replace(find(slide, "What is ROS?"), title)
    replace(find(slide, "A shared framework for building robotics software."), subtitle)
    return slide


def add_photo(slide, path, x, y, w, h):
    """Fill a consistent frame using PowerPoint's editable image crop."""
    from PIL import Image

    with Image.open(path) as im:
        ratio = im.width / im.height
    frame_ratio = w / h
    image = slide.shapes.add_picture(str(path), Pt(x), Pt(y), Pt(w), Pt(h))
    if ratio > frame_ratio:
        image.crop_left = image.crop_right = (1 - frame_ratio / ratio) / 2
    else:
        image.crop_top = image.crop_bottom = (1 - ratio / frame_ratio) / 2
    return image


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def build():
    prs = Presentation(SOURCE)
    original = list(prs.slides)
    credits = json.loads((ASSETS / "credits.json").read_text())

    # Two inserted slides; all other slides retain their original editable shapes.
    about = new_slide(prs, original[1], "About me", "Gilad Lifshitz")
    name = find(about, "Gilad Lifshitz")
    name.top, name.height = Pt(143), Pt(50)
    replace(name, "Gilad Lifshitz", 36)
    box(about, 54, 211, 414, 115)
    text(about, 76, 224, 370, 26, "STUDIES", 16, "blue", True)
    text(about, 76, 257, 370, 65, "Master’s student under Ronen.", 27)
    box(about, 492, 211, 414, 115, "paleTeal")
    text(about, 514, 224, 370, 26, "EXPERIENCE", 16, "teal", True)
    text(about, 514, 257, 370, 65,
         "Robots and drones\nduring army service.", 26)
    box(about, 54, 347, 852, 132, "paleBlue")
    text(about, 76, 359, 808, 24, "YOUR FEEDBACK MATTERS", 16, "blue", True)
    text(about, 76, 388, 808, 84,
         "This is the first time we’re running this course.\n"
         "All the tutorials are new. I’d love your feedback—\n"
         "what works well and what needs improvement.", 24)
    notes(about,
          "Introduce yourself as Gilad Lifshitz, a master’s student under Ronen. "
          "Describe your experience working with robots and drones during army service. "
          "Do not add an unconfirmed surname, degree specialization, or military role.\n\n"
          "This is the first time we’re running this course, and all the tutorials are new. "
          "I’d love your feedback—what works well and what needs improvement. "
          "Invite both positive and critical feedback throughout the semester.")

    agenda = new_slide(prs, original[1], "ROS basics", "Today’s plan")
    for i, label in enumerate([
        "Understand ROS at a high level", "Watch a demonstration",
        "Read the demonstration’s code",
    ]):
        y = 206 + i * 90
        box(agenda, 54, y, 852, 74, "paleTeal" if i == 1 else "white")
        text(agenda, 76, y + 20, 62, 38, f"0{i + 1}", 28, "teal", True)
        text(agenda, 156, y + 20, 724, 40, label, 28, "ink", True)
    notes(agenda,
          "Today we will learn what ROS provides at a high level, watch a live "
          "publisher/subscriber demonstration, and read the Python code that makes it work. "
          "The demonstration and code walkthrough share the dedicated Demo time block. "
          "Assume Python familiarity and little networking knowledge.")

    ids = prs.slides._sldIdLst
    for offset in range(2):
        element = ids[-1]
        ids.remove(element)
        ids.insert(1, element)

    opening = original[0]
    for shape in list(opening.shapes):
        if Pt(100) < shape.top < Pt(490):
            remove(shape)
    replace(find(opening, "WEEK 01 / ROS BASICS"), "WELCOME TO THE COURSE")
    title = find(opening, "ROS basics")
    title.top, title.height = Pt(167), Pt(143)
    replace(title, TITLES[0], 54)
    text(opening, 54, 330, 852, 48, "Week 1", 30, "teal", True)
    notes(opening,
          "Week 1 — Robotics & Autonomous Systems\n"
          "Display the course title while students settle.\n"
          "Audience: Python familiarity; little networking knowledge.\n"
          "The final deck has 15 main slides and five hidden backup slides (16–20).\n"
          "The main session is 45 minutes: explanation/introduction 25; code 7; "
          "run/inspect 6; experiments 4; questions/recap 3.\n"
          "ROS 2 Jazzy is the teaching baseline, not a claim about the newest distribution.")

    definition = original[1]
    replace(find(definition, "What is ROS?"), TITLES[3], 38)
    replace(find(definition, "A shared framework for building robotics software."),
            "Despite its name, ROS is a software framework\n"
            "that runs on an operating system.", 25)

    applications = original[2]
    clear_body(applications)
    replace(find(applications, "Integrate sensing, decisions, and actuation in one system."),
            "One framework. Many kinds of autonomous systems.")
    for i, (label, filename) in enumerate([
        ("Autonomous vehicles", "autonomous_vehicle.jpg"),
        ("Robotics", "robotics.jpg"), ("Drones", "drone.jpg"),
    ]):
        x = 54 + i * 292
        box(applications, x, 208, 268, 258)
        add_photo(applications, ASSETS / filename, x, 208, 268, 198)
        text(applications, x + 8, 426, 252, 34, label, 25,
             "teal" if i == 1 else "blue", True, PP_ALIGN.CENTER)
    credit_notes = "\n\n".join(
        f"{c['title']}\nAuthor: {c['author']}\nLicense: {c['license']}\n"
        f"License URL: {c['license_url']}\nSource: {c['source']}\nChanges: {c['changes']}"
        for c in credits
    )
    notes(applications,
          "Purpose: Show three recognizable fields where ROS is used: autonomous "
          "vehicles, robotics, and drones. Explain the shared need to connect sensors, "
          "perception, planning, and control. These photographs illustrate application "
          "fields; they do not establish that a particular pictured product runs ROS.\n\n"
          "Retained teaching context from the updated deck:\n"
          + applications.notes_slide.notes_text_frame.text
          + "\n\nIMAGE CREDITS\n" + credit_notes)

    replace(find(original[7], "One graph, multiple computers"), "ROS domain")

    communication = original[8]
    replace(find(communication, "ROS 2 architecture and DDS"), TITLES[10], 38)
    replace(find(communication, "DDS = Data Distribution Service. RTPS = its wire protocol."),
            "DDS = Data Distribution Service: discovery and message exchange.", 25)
    find(communication, "Shared domain · type · compatible QoS").top = Pt(426)
    text(communication, 54, 464, 852, 28,
         "Middleware and communication settings can be configured when needed.",
         21, "muted")
    notes(communication, communication.notes_slide.notes_text_frame.text
          + "\n\nCONFIGURATION (speaker reference; no live setup change required):\n"
          "The course setup uses DDS through ROS 2's replaceable RMW interface. "
          "RMW_IMPLEMENTATION selects an installed middleware implementation, for example "
          "rmw_fastrtps_cpp or rmw_cyclonedds_cpp. Configure matching discovery/domain "
          "settings in each terminal and compatible QoS on endpoints. DDS vendor profiles "
          "can tune transports and discovery when needed; profile syntax is vendor-specific. "
          "Do not imply that every ROS 2 deployment must use DDS. DDS handles network "
          "discovery/data exchange; the executor schedules application callbacks.\n"
          "Reference: " + RMW_REFERENCE)

    qos = original[9]
    text(qos, 54, 173, 852, 29,
         "More QoS policies are available, including deadline, lifespan, and liveliness.",
         20, "muted")
    notes(qos, qos.notes_slide.notes_text_frame.text
          + "\n\nADDITIONAL POLICIES: Deadline describes the expected interval between "
          "messages; lifespan limits how long a sample remains valid; liveliness helps "
          "detect whether publishers are still alive. Mention that these exist without "
          "expanding the three-policy overview. Official Jazzy reference: " + QOS_REFERENCE)

    lifecycle = original[10]
    replace(find(lifecycle, "Execution and lifecycle"), "Lifecycle of a node")
    for shape in list(lifecycle.shapes):
        if shape.top >= Pt(190) and shape.top < Pt(490) and shape.left >= Pt(492):
            remove(shape)
    # Enlarge and center the original ordinary-node panel and its three steps.
    for shape in lifecycle.shapes:
        if not Pt(190) <= shape.top < Pt(490):
            continue
        if shape.shape_id == 8:
            shape.left, shape.width = Pt(170), Pt(620)
        elif shape.shape_id == 9:
            shape.left, shape.width = Pt(194), Pt(572)
            replace(shape, shape.text, 26)
            shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        elif shape.shape_id in (10, 13, 16):
            shape.left, shape.width = Pt(194), Pt(572)
        elif shape.shape_id in (11, 14, 17):
            shape.left, shape.width = Pt(206), Pt(548)
            replace(shape, shape.text, 28)
        else:
            shape.left = Pt(480)
    old_lifecycle_notes = lifecycle.notes_slide.notes_text_frame.text
    # Keep the ordinary-runtime explanation; managed lifecycle belongs to backup 19.
    ordinary_notes = old_lifecycle_notes.split("- A managed lifecycle node")[0]
    ordinary_notes = re.split(r"(?im)^.*managed (?:lifecycle|node).*$", ordinary_notes)[0]
    notes(lifecycle,
          "Purpose: Explain the lifecycle of the ordinary Node used in this demo.\n\n"
          "rclpy.init() → construct node/endpoints/timer → rclpy.spin(node)\n"
          "The executor waits for ready work and runs timer/subscription callbacks.\n"
          "Ctrl+C ends spinning; finally destroys the node and shuts down the context.\n"
          "DDS delivers data; the executor runs your callbacks.\n"
          "These are application startup/runtime/cleanup steps, rather than managed "
          "lifecycle states. The optional managed lifecycle diagram is backup slide 19.\n\n"
          "Retained ordinary-node teaching notes:\n" + ordinary_notes)

    replace(find(original[12], "Let's read it. Then run it."), "Demo time")

    for number, slide in enumerate(prs.slides, 1):
        footer = next(s for s in slide.shapes if s.has_text_frame
                      and s.left >= Pt(800) and s.top >= Pt(500))
        replace(footer, f"{number:02d}")
        if number <= 15:
            slide._element.attrib.pop("show", None)
        else:
            slide._element.set("show", "0")
        body = slide.notes_slide.notes_text_frame.text
        body = re.sub(r"^Timing:\s*[\d.]+\s+minutes?\.?\s*", "", body)
        body = re.sub(r"^total\s*—\s*", "Demo block: ", body)
        prefix = f"Final slide {number} — {TITLES[number - 1].replace(chr(10), ' ')}\n"
        prefix += (f"Timing: {MINUTES[number - 1]:g} minutes.\n\n" if number <= 15
                   else "Optional backup — hidden during normal slideshow playback.\n\n")
        notes(slide, prefix + body)

    prs.core_properties.title = "Robotics & Autonomous Systems — Week 1: ROS basics"
    prs.core_properties.subject = "Final Week 1 presentation; 15 main slides and 5 hidden backups"
    prs.core_properties.author = "Gilad Lifshitz"
    prs.save(OUTPUT)
    return validate_powerpoint()


def protected_hashes():
    paths = [ROOT / n for n in (
        "ROS_BASICS.pptx", "ROS_BASICS.pdf", "ROS_BASICS_UPDATED.pptx",
        "ROS_BASICS_UPDATED.pdf", "PRESENTATION_PLAN.md", "TEACHING_NOTES.md",
        "demo/README.md", "demo/publisher.py", "demo/subscriber.py",
    )]
    return {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}


def shape_xml(slide, excluded=()):
    return [etree.tostring(s._element, method="c14n") for s in slide.shapes
            if s.shape_id not in excluded and not (
                s.left >= Pt(800) and s.top >= Pt(500))]


def validate_powerpoint():
    source, final = Presentation(SOURCE), Presentation(OUTPUT)
    assert len(final.slides) == 20
    assert sum(MINUTES) == 45
    assert (final.slide_width, final.slide_height) == (source.slide_width, source.slide_height)
    with ZipFile(OUTPUT) as package:
        assert package.testzip() is None
        for name in package.namelist():
            if name.endswith((".xml", ".rels")):
                etree.fromstring(package.read(name))
    slides = []
    for n, s in enumerate(final.slides, 1):
        assert find(s, TITLES[n - 1]) is not None
        assert (s._element.get("show") == "0") == (n > 15)
        assert len(s.notes_slide.notes_text_frame.text) > 100
        footer = next(t for t in s.shapes if t.has_text_frame
                      and t.left >= Pt(800) and t.top >= Pt(500))
        assert footer.text == f"{n:02d}"
        slides.append({"number": n, "title": TITLES[n - 1].replace("\n", " "),
                       "minutes": MINUTES[n - 1], "hidden": n > 15})

    unchanged = {6: 4, 7: 5, 8: 6, 9: 7, 14: 12,
                 16: 14, 17: 15, 18: 16, 19: 17, 20: 18}
    for final_number, source_number in unchanged.items():
        assert shape_xml(final.slides[final_number - 1]) == shape_xml(
            source.slides[source_number - 1]), f"Changed preserved slide {final_number}"
        if final_number > 15:
            assert source.slides[source_number - 1].notes_slide.notes_text_frame.text in (
                final.slides[final_number - 1].notes_slide.notes_text_frame.text
            )
    # Domain and demo change only their title and footer.
    for final_number, source_number in ((10, 8), (15, 13)):
        assert shape_xml(final.slides[final_number - 1], (3,)) == shape_xml(
            source.slides[source_number - 1], (3,))
    # Original definition layers, communication stack/arrows, and QoS cards survive.
    for fi, si, shape_ids in [(4, 2, range(10, 16)), (11, 9, range(8, 33)),
                              (12, 10, range(8, 23))]:
        a = {s.shape_id: s for s in final.slides[fi - 1].shapes}
        b = {s.shape_id: s for s in source.slides[si - 1].shapes}
        for sid in shape_ids:
            assert etree.tostring(a[sid]._element) == etree.tostring(b[sid]._element)
    runtime = " ".join(s.text for s in final.slides[12].shapes if s.has_text_frame)
    assert "ORDINARY NODE" in runtime and "MANAGED NODE" not in runtime
    assert all(x not in runtime for x in ("Unconfigured", "Inactive", "Finalized"))
    assert len([s for s in final.slides[4].shapes if s.shape_type == 13]) == 3
    return {"slides": slides, "mainSlides": 15, "hiddenBackups": 5,
            "timingMinutes": 45, "preservedSlides": list(unchanged),
            "editableDiagramChecks": "passed", "packageCheck": "passed"}


def export_pdf(soffice):
    with tempfile.TemporaryDirectory(prefix="ros-final-export-") as directory:
        work = Path(directory)
        # Export a temporary all-visible copy, leaving backups hidden in the deliverable.
        deck = Presentation(OUTPUT)
        for slide in deck.slides:
            slide._element.attrib.pop("show", None)
        export_source = work / OUTPUT.name
        deck.save(export_source)
        command = [soffice, f"-env:UserInstallation={(work / 'profile').as_uri()}",
                   "--headless", "--nologo", "--nodefault", "--nofirststartwizard",
                   "--convert-to", "pdf:impress_pdf_Export", "--outdir", str(work),
                   str(export_source)]
        result = subprocess.run(command, capture_output=True, text=True, timeout=180)
        pdf = export_source.with_suffix(".pdf")
        if result.returncode or not pdf.exists():
            raise RuntimeError("LibreOffice PDF export failed:\n" + result.stdout + result.stderr)
        shutil.copyfile(pdf, OUTPUT.with_suffix(".pdf"))


def render_and_validate_pdf():
    import pymupdf

    doc = pymupdf.open(OUTPUT.with_suffix(".pdf"))
    assert len(doc) == 20, f"PDF contains {len(doc)} pages instead of 20"
    preview = SLIDES / "final_preview"
    preview.mkdir(exist_ok=True)
    normalize = lambda s: " ".join(unicodedata.normalize("NFKC", s).split())
    presentation = Presentation(OUTPUT)
    text_frames_checked = 0
    for index, page in enumerate(doc):
        content = normalize(page.get_text())
        assert normalize(TITLES[index]) in content, f"PDF title missing on slide {index + 1}"
        assert f"{index + 1:02d}" in content
        # Verify every text frame survived export, including code and multiline notes.
        compact_content = re.sub(r"\s+", "", content)
        for shape in presentation.slides[index].shapes:
            if not shape.has_text_frame or not shape.text.strip():
                continue
            expected = re.sub(r"\s+", "", normalize(shape.text))
            assert expected in compact_content, (
                f"Missing/truncated text on slide {index + 1}: {shape.text}"
            )
            text_frames_checked += 1
        if index >= 15:
            assert "OPTIONAL / DEEPER DISCUSSION" in content
        page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(
            preview / f"slide_{index + 1:02d}.png")
    assert "Ronen" in doc[1].get_text()
    assert "feedback" in doc[1].get_text()
    assert all(word in doc[11].get_text() for word in ("deadline", "lifespan", "liveliness"))
    assert "MANAGED NODE" not in doc[12].get_text()
    assert all(word in doc[14].get_text() for word in (
        "/greeting_publisher", "/chatter", "/greeting_subscriber", "std_msgs/msg/String"))
    result = {"pages": len(doc), "allTitlesAndFooters": "passed",
              "textFramesChecked": text_frames_checked, "allSlideTextPresent": True,
              "backupPagesIncluded": True, "requiredContent": "passed",
              "previewDirectory": str(preview.relative_to(ROOT))}
    doc.close()
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--soffice", default=shutil.which("libreoffice") or shutil.which("soffice"))
    parser.add_argument("--no-pdf", action="store_true", help="Build and validate only the PPTX")
    args = parser.parse_args()
    if not args.no_pdf and not args.soffice:
        parser.error("LibreOffice is required: install it or pass --soffice PATH")
    before = protected_hashes()
    report = build()
    if not args.no_pdf:
        export_pdf(args.soffice)
        report["pdf"] = render_and_validate_pdf()
    assert protected_hashes() == before, "An original deck or protected teaching file changed"
    report["protectedFiles"] = before
    report["originalFilesUnchanged"] = True
    (SLIDES / "final_validation.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(f"Created {OUTPUT}")
    if not args.no_pdf:
        print(f"Created {OUTPUT.with_suffix('.pdf')}")
    print("Validated: 20 slides, 15 main slides, 5 hidden backups; originals unchanged.")


if __name__ == "__main__":
    main()
